import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

with open("results.json") as f:
    data = json.load(f)

order = {"WATCH": 0, "BUY": 1, "SELL": 2}


def sig(s):
    return s["signal"].split()[0]


stocks = sorted(data["stocks"], key=lambda s: (order.get(sig(s), 3), s["symbol"]))

wb = Workbook()
ws = wb.active
ws.title = "Scopevision Screener"

# Row 1: title, Row 2: last updated, Row 3: table header
ws.append(["Scopevision Screener"])
ws.append(["Last updated: " + data["last_updated"][:16].replace("T", " ") + " IST"])

header = ["Symbol", "LTP", "Point", "VHIGH", "VLOW", "Signal", "Short T", "Long T", "SL"]
ws.append(header)

for s in stocks:
    ws.append([
        s["symbol"], s["ltp"], s["poc"], s["vah"], s["val"], sig(s),
        s["short_target"], s["long_target"], s["stop_loss"],
    ])

ws.merge_cells("A1:I1")
ws["A1"].font = Font(size=14, bold=True)
ws["A1"].alignment = Alignment(horizontal="left")
ws["A2"].font = Font(italic=True, color="6B7280")

head_fill = PatternFill("solid", fgColor="1F2937")
for cell in ws[3]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = head_fill
    cell.alignment = Alignment(horizontal="center")

colors = {"BUY": "15803D", "SELL": "B91C1C", "WATCH": "B45309"}
for row in ws.iter_rows(min_row=4):
    c = colors.get(row[5].value)
    if c:
        row[5].font = Font(bold=True, color=c)

for col, width in zip("ABCDEFGHI", [16, 10, 10, 10, 10, 10, 11, 11, 11]):
    ws.column_dimensions[col].width = width

ws.freeze_panes = "A4"
ws.auto_filter.ref = f"A3:I{ws.max_row}"

wb.save("Screener_report.xlsx")
print("Excel ready")
